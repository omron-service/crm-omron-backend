def generate_inventory_excel(headers, rows, sheet_title="Laporan"):
    """
    FUNGSI BARU - tambahkan ini ke akhir file app/services/excel_export.py Anda.
    Tidak menyentuh/mengubah fungsi generate_service_report_excel yang sudah ada.

    Membuat file Excel generik dari daftar header kolom + baris data, dipakai
    untuk semua laporan Inventory Part (stok list, mutasi, stok opname).

    headers: list[str]  - judul kolom, mis. ["Kode Sparepart", "Nama Sparepart", ...]
    rows:    list[list]  - tiap item adalah satu baris, urutan nilai sesuai headers
    """
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "Laporan")[:31]  # batas judul sheet Excel = 31 karakter

    ws.append(headers)
    header_fill = PatternFill(start_color="0056B3", end_color="0056B3", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in rows:
        ws.append(row)

    for i, header in enumerate(headers, start=1):
        col_letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col_letter].width = max(15, len(str(header)) + 2)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
