"""
Generator file Excel "FORMULIR PERBAIKAN" (Service Report) per tiket, meniru
persis layout template asli yang diberikan (ServiceReport-OEC-....xlsx):
judul, data pelanggan/alat dalam kotak berlabel, tabel sparepart, syarat &
ketentuan, dan logo Omron.
"""
import io
import os
from datetime import datetime
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils.cell import range_boundaries

ASSETS_DIR = os.path.join(os.path.dirname(__file__), "..", "assets")
LOGO_PATH = os.path.join(ASSETS_DIR, "omron_logo.png")
BANNER_PATH = os.path.join(ASSETS_DIR, "omron_healthcare_banner.png")


def _mask_phone(phone: Optional[str]) -> str:
    """
    Sensor nomor HP/WA untuk file Service Report yang di-download - 4 digit
    awal dan 2 digit akhir tetap tampil, sisanya diganti 'x'.
    Contoh: 081212345678 -> 0812xxxxxx78
    """
    if not phone:
        return "-"
    digits_only = phone.strip()
    if len(digits_only) <= 6:
        return digits_only  # terlalu pendek utk disensor tanpa jadi tidak berguna
    prefix = digits_only[:4]
    suffix = digits_only[-2:]
    masked_len = len(digits_only) - 4 - 2
    return f"{prefix}{'x' * masked_len}{suffix}"

HEADER_FILL = PatternFill(start_color="FFF6F8FC", end_color="FFF6F8FC", fill_type="solid")
THIN = Side(style="thin", color="000000")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

FONT_TITLE = Font(name="Calibri", size=18, bold=True)
FONT_LABEL = Font(name="Calibri", size=11, bold=False)
FONT_LABEL_BOLD = Font(name="Calibri", size=11, bold=True)
FONT_FOOTER = Font(name="Calibri", size=14, bold=True, color="FF005EB8")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

COLUMN_WIDTHS = {"A": 3.73, "B": 21.18, "C": 23.0, "D": 14.18, "E": 15.18, "F": 15.27, "G": 8.82}
ROW_HEIGHTS = {
    1: 12, 2: 12, 3: 19.25, 4: 12, 5: 12, 6: 12, 7: 12, 8: 12, 9: 12, 10: 24,
    11: 12, 12: 12, 13: 24, 14: 100, 15: 12, 16: 12, 17: 12, 18: 12, 19: 12,
    20: 12, 21: 12, 22: 12, 23: 12, 24: 12, 25: 12, 26: 12, 27: 12, 28: 12,
    29: 12, 30: 12, 31: 12, 32: 12, 33: 12, 34: 12, 35: 12, 36: 12, 37: 12,
    38: 12, 39: 12, 40: 17.65,
}

TERMS_1 = (
    "1. Garansi tidak berlaku jika kerusakan disebabkan oleh kesalahan penggunaan "
    "atau efek lingkungan (hama, serangga, terkena cairan, robek, terbakar, dll)."
)
TERMS_2 = (
    "2. Omron Healthcare Service Center tidak bertanggung jawab atas kehilangan dan "
    "kerusakan alat yang disebabkan oleh FORCE MAJEURE (gempa bumi, banjir, kebakaran, "
    "pencurian, perampokan, dll)."
)


def _fmt_date(value) -> str:
    if not value:
        return "-"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _block(ws, rng: str, value=None, font=FONT_LABEL, fill=None, align=ALIGN_LEFT, border=BOX):
    """
    Mengisi & memformat satu "blok" sel - baik itu sel tunggal (mis. "B6") atau
    range gabungan (mis. "E7:F8") - dan yang PALING PENTING: menerapkan border
    ke SEMUA sel di dalam range tersebut, bukan cuma sel pojok kiri-atas.

    Ini memperbaiki masalah "border tidak rapi": kalau border cuma diset di
    satu sel pada range yang di-merge, Excel menampilkan kotak yang terlihat
    "putus" di sisi kanan/bawahnya karena sel-sel lain di dalam range itu
    tidak punya border sama sekali.
    """
    if ":" in rng:
        ws.merge_cells(rng)
    min_col, min_row, max_col, max_row = range_boundaries(rng)

    if value is not None:
        ws.cell(row=min_row, column=min_col).value = value

    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.font = font
            cell.alignment = align
            if fill:
                cell.fill = fill


def generate_ticket_service_report(ticket, spareparts: Optional[Iterable] = None) -> io.BytesIO:
    """
    ticket: instance ServiceTicket (butuh atribut: ticket_number, warranty_status,
        customer_name, received_date, completed_date, customer_phone,
        customer_address, device_model, serial_number, accessories, complaint,
        technician_analysis, created_at).
    spareparts: iterable TicketSparePart (name, code, quantity, price, slot_no)
        - maksimal 5 baris ditampilkan sesuai template asli.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "BizForm"

    for col, w in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = w
    for row, h in ROW_HEIGHTS.items():
        ws.row_dimensions[row].height = h

    # ---- Pengaturan halaman cetak: paksa semua kolom (A-G) muat dalam 1
    # halaman lebar, supaya tidak ada kolom yang terpotong saat dibuka/dicetak
    # (mis. logo, kolom Harga, No. Serial yang sebelumnya kepotong di tepi).
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0  # 0 = tinggi menyesuaikan, tidak dipaksa 1 halaman
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:G40"

    # ---- Judul + logo ----
    ws.merge_cells("B2:D3")
    title_cell = ws["B2"]
    title_cell.value = "FORMULIR  PERBAIKAN"
    title_cell.font = FONT_TITLE
    title_cell.alignment = ALIGN_LEFT

    if os.path.exists(LOGO_PATH):
        # Dibungkus try/except: penyisipan gambar butuh library Pillow. Kalau
        # suatu saat Pillow belum/tidak ter-install di server, laporan TETAP
        # berhasil dibuat (tanpa logo) daripada gagal total (500 error).
        try:
            logo_img = XLImage(LOGO_PATH)
            logo_img.width = 190
            logo_img.height = 50
            ws.add_image(logo_img, "E2")
        except ImportError:
            pass

    # ---- Data tiket & pelanggan ----
    # Setiap baris berikut: LABEL (kotak abu-abu, rata kiri) + VALUE (kotak
    # putih). Border SELALU diterapkan ke seluruh range (lihat _block), jadi
    # baik sel tunggal maupun sel gabungan sama-sama rapi.
    _block(ws, "B6", "NO - ERF Number", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "C6", ticket.ticket_number, align=ALIGN_CENTER)
    _block(ws, "D6", "KONTRAK ID", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "E6:F6", "", align=ALIGN_CENTER)

    _block(ws, "B7", "STATUS GARANSI", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "C7", ticket.warranty_status or "-", align=ALIGN_CENTER)
    _block(ws, "D7:D8", "NAMA PEMILIK", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "E7:F8", ticket.customer_name or "-", align=ALIGN_CENTER)

    _block(ws, "B8", "TANGGAL TERIMA", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "C8", _fmt_date(ticket.received_date or ticket.created_at), align=ALIGN_CENTER)

    _block(ws, "B9", "TANGGAL SELESAI", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "C9", _fmt_date(ticket.completed_date), align=ALIGN_CENTER)
    _block(ws, "D9", "NO. TELEPON", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "E9:F9", _mask_phone(ticket.customer_phone), align=ALIGN_CENTER)

    _block(ws, "B10", "ALAMAT", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "C10:F10", ticket.customer_address or "-", align=ALIGN_LEFT)

    _block(ws, "B11", "PRODUK", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "C11", ticket.device_model or "-", align=ALIGN_CENTER)
    _block(ws, "D11", "NO. SERIAL", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "E11:F11", ticket.serial_number or "-", align=ALIGN_CENTER)

    _block(ws, "B12", "AKSESORIS", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "C12:F12", ticket.accessories or "-", align=ALIGN_LEFT)

    _block(ws, "B13", "PROBLEM", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "C13:F13", ticket.complaint or "-", align=ALIGN_LEFT_WRAP)

    _block(ws, "B14", "KESIMPULAN", fill=HEADER_FILL, align=ALIGN_LEFT)
    _block(ws, "C14:F14", ticket.technician_analysis or "-", align=ALIGN_LEFT_WRAP)

    # ---- Tabel sparepart (maks 5 baris, sesuai template asli) ----
    _block(ws, "B16", "Nama Spare Part", font=FONT_LABEL_BOLD, fill=HEADER_FILL, align=ALIGN_CENTER)
    _block(ws, "C16", "Kode Spare Part", font=FONT_LABEL_BOLD, fill=HEADER_FILL, align=ALIGN_CENTER)
    _block(ws, "D16", "Jumlah", font=FONT_LABEL_BOLD, fill=HEADER_FILL, align=ALIGN_CENTER)
    _block(ws, "E16:F16", "Harga. Rp.", font=FONT_LABEL_BOLD, fill=HEADER_FILL, align=ALIGN_CENTER)

    sp_list = list(spareparts or [])[:5]
    total = 0.0
    for i in range(5):
        row = 17 + i
        sp = sp_list[i] if i < len(sp_list) else None
        _block(ws, f"B{row}", sp.name if sp and sp.name else "", align=ALIGN_LEFT)
        _block(ws, f"C{row}", sp.code if sp and sp.code else "", align=ALIGN_CENTER)
        _block(ws, f"D{row}", sp.quantity if sp and sp.quantity else "", align=ALIGN_CENTER)
        line_total = (sp.price or 0) * (sp.quantity or 0) if sp else 0
        _block(ws, f"E{row}:F{row}", line_total if sp else "", align=ALIGN_RIGHT)
        total += line_total

    _block(ws, "D22", "Total", border=Border(), align=ALIGN_RIGHT)
    _block(ws, "E22:F22", total, align=ALIGN_RIGHT)

    # ---- Syarat & Ketentuan (tanpa border, sesuai template asli) ----
    ws["B25"].value = "Syarat & Ketentuan"
    ws["B25"].font = FONT_LABEL_BOLD

    ws.merge_cells("B26:F27")
    ws["B26"].value = TERMS_1
    ws["B26"].font = FONT_LABEL
    ws["B26"].alignment = ALIGN_LEFT_WRAP

    ws.merge_cells("B28:F30")
    ws["B28"].value = TERMS_2
    ws["B28"].font = FONT_LABEL
    ws["B28"].alignment = ALIGN_LEFT_WRAP

    ws["B35"].value = "Pemilik : "
    ws["B35"].font = FONT_LABEL
    ws["D35"].value = "Teknisi:"
    ws["D35"].font = FONT_LABEL
    ws["D35"].alignment = ALIGN_CENTER

    if os.path.exists(BANNER_PATH):
        try:
            banner_img = XLImage(BANNER_PATH)
            banner_img.width = 209
            banner_img.height = 74
            ws.add_image(banner_img, "D39")
        except ImportError:
            pass

    ws["B40"].value = "ALL for Healthcare"
    ws["B40"].font = FONT_FOOTER

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
