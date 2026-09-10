import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO

def generate_service_report_excel(rows):
    """
    Laporan Excel untuk tombol "Download Excel" di menu Data Service
    (Pusat/Cabang/Pickup Center) - format PERSIS sesuai template yang diberikan
    (Form_donwload_tiket_service_v1.xlsx), 41 kolom termasuk 3 kolom "NA" yang
    sengaja dikosongkan (spacer, sesuai template asli).

    `rows` adalah list of dict, satu dict per tiket, dengan key-key berikut
    (dibangun di app/main.py dari data ServiceTicket + TicketSparePart):
    ticket_number, nama_teknisi, received_date, completed_date, customer_name,
    province, city, customer_address, instansi_name, customer_phone,
    customer_phone_2, product_category, device_model, serial_number,
    warranty_period, warranty_status, accessories, complaint,
    technician_analysis, symptom_code, product_origin, leadtime_days,
    remarks, status, notes, total_price,
    sp1_name, sp1_qty, sp1_code, sp1_price,
    sp2_name, sp2_qty, sp2_code, sp2_price,
    sp3_name, sp3_qty, sp3_code, sp3_price.
    """
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    HEADERS = [
        "No. Tiket", "NA", "Nama Teknisi", "Tangal Diterima", "Tanggal Selesai",
        "Nama Pemilik", "Provinsi", "Kota/ Kabupaten", "Alamat", "Nama Instansi",
        "No. HP/ Whatsapp 1", "No. HP/ Whatsapp 2", "Produk Kategori", "Model Alat",
        "Serial No. Alat", "Warranty Period", "Status Garansi", "Aksesoris",
        "Keluhan Pelanggan", "NA", "Analisa Teknisi", "Symptom", "Asal produk",
        "Lead Time", "Remarks", "Repair Status", "Catatan", "NA", "Total Price",
        "Name of Spare Part #1", "Number of Spare Part #1", "Spare Part Code #1", "Price #1",
        "Name of Spare Part #2", "Number of Spare Part #2", "Spare Part Code #2", "Price #2",
        "Name of Spare Part #3", "Number of Spare Part #3", "Spare Part Code #3", "Price #3",
    ]

    # Urutan key HARUS PERSIS selaras dengan HEADERS di atas (termasuk 3 kolom
    # "NA" yang memang sengaja dikosongkan - diberi key unik _na1/_na2/_na3
    # supaya tidak bentrok satu sama lain di dict Python).
    FIELD_ORDER = [
        "ticket_number", "_na1", "nama_teknisi", "received_date", "completed_date",
        "customer_name", "province", "city", "customer_address", "instansi_name",
        "customer_phone", "customer_phone_2", "product_category", "device_model",
        "serial_number", "warranty_period", "warranty_status", "accessories",
        "complaint", "_na2", "technician_analysis", "symptom_code", "product_origin",
        "leadtime_days", "remarks", "status", "notes", "_na3", "total_price",
        "sp1_name", "sp1_qty", "sp1_code", "sp1_price",
        "sp2_name", "sp2_qty", "sp2_code", "sp2_price",
        "sp3_name", "sp3_qty", "sp3_code", "sp3_price",
    ]

    COLUMN_WIDTHS = [12, 4, 16, 13, 13, 18, 13, 15, 22, 16, 15, 15, 14, 14, 14, 13,
                     14, 16, 25, 4, 20, 12, 13, 9, 22, 14, 20, 4, 12,
                     18, 12, 14, 10, 18, 12, 14, 10, 18, 12, 14, 10]

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Service"

    # Landscape + fit-to-width supaya kalau dicetak/di-print-preview tidak
    # terpotong jadi berhalaman-halaman (41 kolom cukup lebar).
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0056B3", end_color="0056B3", fill_type="solid")
    for col_idx, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width
    ws.freeze_panes = "A2"

    for row_data in rows:
        ws.append([row_data.get(key, "") if not key.startswith("_na") else "" for key in FIELD_ORDER])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
def generate_payment_status_excel(rows):
    """
    Laporan 'Status Payment Service' - 22 kolom PERSIS sesuai template resmi
    yang diberikan (contoh_laporan_payment_status.xls).

    `rows` adalah list of dict dgn key: no, ticket_number, owner_name, email,
    phone, address, id_number, warranty_status, harga, ppn, total, pph23,
    admin_bank, ppn_for_doku, status_repair, invoice_number, invoice_date,
    payment_channel, payment_status, created_at, paid_at, departement.
    """
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    HEADERS = [
        "No", "Nomer Tiket", "Nama Pemilik", "Email Pemilik", "No. HP/ Whatsapp 1",
        "Alamat", "NIK/ NPWP", "Status Garansi", "Harga", "PPN", "Total",
        "PPH23", "Biaya Admin", "PPN for DOKU", "Status Repair", "No. Invoice",
        "Tanggal Invoice", "Channel Pembayaran", "Status Pembayaran", "Dibuat pada",
        "Tanggal Bayar", "Departement",
    ]
    FIELD_ORDER = [
        "no", "ticket_number", "owner_name", "email", "phone", "address", "id_number",
        "warranty_status", "harga", "ppn", "total", "pph23", "admin_bank", "ppn_for_doku",
        "status_repair", "invoice_number", "invoice_date", "payment_channel", "payment_status",
        "created_at", "paid_at", "departement",
    ]
    COLUMN_WIDTHS = [5, 16, 18, 22, 15, 22, 16, 15, 14, 12, 14, 10, 12, 12, 16, 16, 16, 16, 15, 16, 16, 14]

    wb = Workbook()
    ws = wb.active
    ws.title = "Status Payment"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0056B3", end_color="0056B3", fill_type="solid")
    for col_idx, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width
    ws.freeze_panes = "A2"

    for row_data in rows:
        ws.append([row_data.get(key, "") for key in FIELD_ORDER])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
    
def generate_inventory_excel(headers, rows, sheet_title="Laporan"):
    """
    FUNGSI BARU - tambahkan ini ke akhir file app/services/excel_export.py Anda.
    Tidak menyentuh/mengubah fungsi generate_service_report_excel yang sudah ada.

    Membuat file Excel generik dari daftar header kolom + baris data, dipakai
    untuk semua laporan Inventory Part (stok list, mutasi, stok opname).

    headers: list[str]  - judul kolom, mis. ["Kode Sparepart", "Nama Sparepart", ...]
    rows:    list[list]  - tiap item adalah satu baris, urutan nilai sesuai headers
    """
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "Laporan")[:31]  # batas judul sheet Excel = 31 karakter

    ws.append(headers)
    header_fill = PatternFill(start_color="0056B3", end_color="0056B3", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in rows:
        ws.append(row)

    for i, header in enumerate(headers, start=1):
        col_letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col_letter].width = max(15, len(str(header)) + 2)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
