import logging
from datetime import datetime, timedelta
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy import desc
from sqlalchemy.orm import Session
from pydantic import BaseModel, Field

from app.db.session import get_db
from app.core.deps import get_current_user, require_department_access
from app.models.schema import PartCatalog, PartStock, PartStockMovement, PartStockOpname, User, BranchCatalog
from app.services.excel_export import generate_inventory_excel

logger = logging.getLogger("uvicorn.error")

router = APIRouter(prefix="/api/v1/inventory-parts", tags=["Inventory Part"])

VALID_INV_LOCATIONS = ("pusat", "cabang")

# Jenis pergerakan yang boleh DIBUAT MANUAL oleh staff (tombol + / upload Excel)
# di masing-masing lokasi, dan arahnya (+1 = item masuk, -1 = item keluar).
# CATATAN: "terpakai_pusat"/"terpakai_cabang" SENGAJA tidak ada di sini lagi
# (tombol manualnya dihapus) - tapi "terpakai_pusat" TETAP otomatis tercatat
# lewat sparepart tiket Pickup Center (lihat _deduct_pusat_stock_for_pickup di
# services_api.py, jalur kode terpisah yang tidak lewat validasi dict ini).
MOVEMENT_RULES = {
    "pusat": {
        "terima_gudang": 1,        # Terima sparepart dari Gudang (item masuk)
        "kirim_ke_cabang": -1,     # Kirim sparepart ke Cabang (item keluar)
        "terima_dari_cabang": 1,   # Terima sparepart dari Cabang (item masuk)
    },
    "cabang": {
        "terima_dari_pusat": 1,      # Terima sparepart dari Pusat (item masuk)
        "kirim_balik_ke_pusat": -1,  # Kirim balik sparepart ke Pusat (item keluar)
    },
}

# Jenis pergerakan yang BOLEH DILAPORKAN (buat download Excel) per lokasi - beda
# dari MOVEMENT_RULES di atas karena "terpakai_pusat" tetap harus bisa dilaporkan
# meski tidak lagi bisa dibuat manual (datanya otomatis dari tiket Pickup Center).
REPORTABLE_MOVEMENT_TYPES = {
    "pusat": {**MOVEMENT_RULES["pusat"], "terpakai_pusat": -1},
    "cabang": {**MOVEMENT_RULES["cabang"], "terpakai_cabang": -1},
}

MOVEMENT_LABELS = {
    "terima_gudang": "Terima dari Gudang",
    "kirim_ke_cabang": "Kirim ke Cabang",
    "terima_dari_cabang": "Terima dari Cabang",
    "terpakai_pusat": "Terpakai di Pusat",
    "terima_dari_pusat": "Terima dari Pusat",
    "kirim_balik_ke_pusat": "Kirim Balik ke Pusat",
    "terpakai_cabang": "Terpakai di Cabang",
}

# Jenis pergerakan yang WAJIB menyertakan nama Cabang (asal/tujuan) - karena
# sekarang ada banyak cabang (lihat menu Kelola Cabang), jadi harus jelas
# cabang MANA yang jadi asal/tujuan pengiriman part, bukan cuma "cabang" secara umum.
# SEMUA 4 jenis pergerakan yang menyentuh Cabang kini wajib - termasuk 2 yang
# dibuat LANGSUNG oleh staff Cabang dari tab mereka sendiri (bukan hasil
# mirror dari Pusat), karena staff Cabang juga tidak terikat ke 1 nama cabang
# spesifik di akun mereka - jadi harus pilih manual setiap kali.
BRANCH_REQUIRED_TYPES = {"kirim_ke_cabang", "terima_dari_cabang", "terima_dari_pusat", "kirim_balik_ke_pusat"}

# Label field "Nama Cabang" yang sesuai utk masing-masing dari 4 jenis pergerakan.
MOVEMENT_BRANCH_LABEL = {
    "kirim_ke_cabang": "Cabang Tujuan",
    "terima_dari_cabang": "Cabang Asal",
    "terima_dari_pusat": "Nama Cabang (cabang Anda)",
    "kirim_balik_ke_pusat": "Nama Cabang (cabang Anda)",
}

# Nilai "Status" sparepart yang valid (sesuai keputusan bisnis).
VALID_PART_STATUS = {"active", "discontinue"}

# PENTING - PASANGAN PERGERAKAN DUA ARAH:
# "Kirim ke Cabang" dan "Terima dari Cabang" di sisi PUSAT adalah PERPINDAHAN
# barang antar lokasi, BUKAN cuma catatan sepihak. Jadi begitu Pusat mengirim
# ke Cabang, stok Pusat berkurang DAN stok Cabang otomatis bertambah (dan
# sebaliknya untuk terima dari cabang) - dicatat sebagai satu transaksi yang
# sama, bukan dua entri manual terpisah yang rawan lupa/selisih.
MIRROR_MOVEMENT_TYPE = {
    "kirim_ke_cabang": "terima_dari_pusat",      # Pusat -1 -> Cabang +1 (otomatis)
    "terima_dari_cabang": "kirim_balik_ke_pusat",  # Pusat +1 -> Cabang -1 (otomatis)
}


class MovementIn(BaseModel):
    location: str
    movement_type: str
    code: str
    name: Optional[str] = None
    quantity: int = Field(..., gt=0)
    note: Optional[str] = None
    device_model: Optional[str] = None       # "Model Alat"
    part_status: Optional[str] = None        # "Status"
    related_branch: Optional[str] = None     # "Nama Cabang Asal"/"Nama Cabang Tujuan"


class OpnameIn(BaseModel):
    location: str
    code: str
    name: Optional[str] = None
    counted_quantity: int = Field(..., ge=0)
    note: Optional[str] = None
    branch_name: Optional[str] = None  # WAJIB kalau location="cabang"


def _get_or_create_part(db: Session, code: str, name: Optional[str], part_cache: Optional[dict] = None) -> PartCatalog:
    if part_cache is not None and code in part_cache:
        part = part_cache[code]
        if name and part.name != name:
            part.name = name
        return part

    part = db.query(PartCatalog).filter(PartCatalog.code == code).first()
    if part is None:
        part = PartCatalog(code=code, name=name)
        db.add(part)
        db.flush()
    elif name and part.name != name:
        part.name = name  # nama boleh diperbarui, kode tetap jadi kunci

    if part_cache is not None:
        part_cache[code] = part
    return part


def _get_or_create_stock_row(db: Session, part_id: int, location: str, branch_name: str = "", stock_cache: Optional[dict] = None) -> PartStock:
    # branch_name HANYA relevan utk location="cabang" - "pusat" selalu "" (satu kolam).
    effective_branch = (branch_name or "").strip() if location == "cabang" else ""
    cache_key = (part_id, location, effective_branch)
    if stock_cache is not None and cache_key in stock_cache:
        return stock_cache[cache_key]

    stock = (
        db.query(PartStock)
        .filter(PartStock.part_id == part_id, PartStock.location == location, PartStock.branch_name == effective_branch)
        .with_for_update()
        .first()
    )
    if stock is None:
        stock = PartStock(part_id=part_id, location=location, branch_name=effective_branch, quantity=0)
        db.add(stock)
        db.flush()

    if stock_cache is not None:
        stock_cache[cache_key] = stock
    return stock


@router.get("/stock/{location}")
def get_stock_list(
    location: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    loc = location.lower().strip()
    if loc not in VALID_INV_LOCATIONS:
        raise HTTPException(status_code=400, detail=f"location harus salah satu dari: {', '.join(VALID_INV_LOCATIONS)}.")
    require_department_access(current_user, loc)

    rows = (
        db.query(PartStock, PartCatalog)
        .join(PartCatalog, PartStock.part_id == PartCatalog.id)
        .filter(PartStock.location == loc)
        .order_by(PartCatalog.name, PartStock.branch_name)
        .all()
    )
    return [
        {
            "part_id": part.id,
            "code": part.code,
            "name": part.name,
            "unit_price": part.unit_price,
            "status": part.status,
            "model_alat": part.model_alat,
            "branch_name": stock.branch_name or None,  # None/kosong utk lokasi "pusat" (1 kolam)
            "quantity": stock.quantity,
            "updated_at": stock.updated_at,
        }
        for stock, part in rows
    ]


def _apply_single_movement(
    db: Session, loc: str, movement_type: str, code: str, name: Optional[str], quantity: int,
    note: Optional[str], user_id: int,
    device_model: Optional[str] = None, part_status: Optional[str] = None, related_branch: Optional[str] = None,
    branch_cache: Optional[dict] = None, part_cache: Optional[dict] = None, stock_cache: Optional[dict] = None,
):
    """
    Logika inti SATU baris pergerakan stok - dipakai bersama oleh endpoint
    manual (/movement) MAUPUN upload massal Excel (/movement/bulk-upload).

    PENTING soal performa & keamanan data (upload file besar seperti 164 baris):
    Fungsi ini SENGAJA didesain dua tahap - VALIDASI DULU SEMUANYA (termasuk
    kecukupan stok di KEDUA lokasi untuk pergerakan berpasangan Pusat<->Cabang),
    baru kemudian benar-benar MENGUBAH ANGKA STOK. Dengan begitu, begitu proses
    "menulis" dimulai, dijamin tidak akan gagal di tengah jalan - jadi TIDAK
    PERLU pakai SAVEPOINT per baris (yang tadinya dipakai untuk jaga-jaga rollback
    per baris, tapi ternyata menyebabkan PostgreSQL sangat lambat kalau dipakai
    lebih dari ~64 kali dalam satu transaksi - persis kasus file besar Anda).
    """
    allowed = MOVEMENT_RULES.get(loc, {})
    if movement_type not in allowed:
        raise ValueError(
            f"movement_type '{movement_type}' tidak berlaku untuk lokasi '{loc}'. "
            f"Pilihan yang valid: {', '.join(allowed.keys())}."
        )
    if not code or not code.strip():
        raise ValueError("Kode Sparepart wajib diisi.")
    if not quantity or quantity <= 0:
        raise ValueError("Jumlah harus lebih dari 0.")

    if part_status and part_status.strip().lower() not in VALID_PART_STATUS:
        raise ValueError(f"Status '{part_status}' tidak valid. Harus 'Active' atau 'Discontinue'.")

    if movement_type in BRANCH_REQUIRED_TYPES:
        if not (related_branch and related_branch.strip()):
            label = MOVEMENT_BRANCH_LABEL.get(movement_type, "Nama Cabang")
            raise ValueError(f"{label} wajib diisi untuk pergerakan '{MOVEMENT_LABELS.get(movement_type, movement_type)}'.")

        # Nama cabang WAJIB cocok dengan daftar aktif di menu Kelola Cabang -
        # mencegah salah ketik nama cabang (mis. "Mendan" vs "Medan") yang
        # kalau dibiarkan bebas teks akan sulit dilacak & direkap nanti.
        # Kalau branch_cache disediakan (dipakai upload massal), lookup nama
        # cabang lewat dict di memori - menghindari 1 query database per BARIS
        # yang untuk file besar (ratusan baris) jadi penyumbang utama lambatnya
        # proses dan risiko timeout ("upstream error").
        branch_key = related_branch.strip().lower()
        if branch_cache is not None:
            official_name = branch_cache.get(branch_key)
        else:
            branch_match = (
                db.query(BranchCatalog)
                .filter(BranchCatalog.name.ilike(related_branch.strip()), BranchCatalog.is_active.is_(True))
                .first()
            )
            official_name = branch_match.name if branch_match else None
        if not official_name:
            raise ValueError(
                f"Nama cabang '{related_branch}' tidak ditemukan di daftar Kelola Cabang (atau sedang nonaktif). "
                f"Pastikan nama cabang persis sama dengan yang terdaftar."
            )
        related_branch = official_name  # normalisasi ke ejaan resmi yang terdaftar

    direction = allowed[movement_type]
    part = _get_or_create_part(db, code.strip(), name, part_cache=part_cache)
    # related_branch diteruskan apa adanya - _get_or_create_stock_row otomatis
    # mengabaikannya kalau loc="pusat" (selalu 1 kolam), dan memakainya sbg
    # kunci cabang spesifik kalau loc="cabang".
    stock = _get_or_create_stock_row(db, part.id, loc, branch_name=related_branch or "", stock_cache=stock_cache)

    # ---- TAHAP 1: VALIDASI DULU (belum ada angka stok yang diubah) ----
    new_quantity = stock.quantity + (direction * quantity)
    if new_quantity < 0:
        raise ValueError(f"Stok tidak mencukupi untuk '{part.code}'. Stok saat ini {stock.quantity}, tidak bisa mengurangi {quantity}.")

    mirror_type = MIRROR_MOVEMENT_TYPE.get(movement_type)
    mirror_stock = None
    mirror_new_quantity = None
    if mirror_type and loc == "pusat":
        mirror_direction = MOVEMENT_RULES["cabang"][mirror_type]
        # Mirror di sisi Cabang HARUS ke cabang SPESIFIK yang sama dgn yang
        # dipilih di form Pusat (related_branch) - bukan lagi kolam gabungan.
        mirror_stock = _get_or_create_stock_row(db, part.id, "cabang", branch_name=related_branch or "", stock_cache=stock_cache)
        mirror_new_quantity = mirror_stock.quantity + (mirror_direction * quantity)
        if mirror_new_quantity < 0:
            raise ValueError(
                f"Stok Cabang tidak mencukupi untuk '{part.code}' pada pergerakan pasangan "
                f"'{MOVEMENT_LABELS.get(mirror_type, mirror_type)}'. Stok Cabang saat ini "
                f"{mirror_stock.quantity}, tidak bisa mengurangi {quantity}."
            )

    # ---- TAHAP 2: SEMUA VALIDASI LOLOS - baru sekarang benar-benar ditulis ----
    normalized_status = None
    if part_status and part_status.strip():
        normalized_status = "Active" if part_status.strip().lower() == "active" else "Discontinue"
    clean_device_model = device_model.strip() if device_model else None
    clean_related_branch = related_branch.strip() if related_branch else None

    stock.quantity = new_quantity
    db.add(PartStockMovement(
        part_id=part.id, location=loc, movement_type=movement_type,
        quantity=quantity, note=note, performed_by_user_id=user_id,
        device_model=clean_device_model, part_status=normalized_status,
        related_branch=clean_related_branch,
    ))
    if normalized_status:
        part.status = normalized_status
    if clean_device_model:
        part.model_alat = clean_device_model

    if mirror_stock is not None:
        mirror_stock.quantity = mirror_new_quantity
        db.add(PartStockMovement(
            part_id=part.id, location="cabang", movement_type=mirror_type,
            quantity=quantity,
            note=(f"Otomatis - pasangan dari '{MOVEMENT_LABELS.get(movement_type, movement_type)}' di Pusat"
                  + (f" ({note})" if note else "")),
            performed_by_user_id=user_id,
            device_model=clean_device_model, part_status=normalized_status,
            related_branch=clean_related_branch,
        ))

    return part, stock


@router.post("/movement")
def create_movement(
    data: MovementIn,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    loc = data.location.lower().strip()
    if loc not in VALID_INV_LOCATIONS:
        raise HTTPException(status_code=400, detail=f"location harus salah satu dari: {', '.join(VALID_INV_LOCATIONS)}.")
    require_department_access(current_user, loc)

    try:
        part, stock = _apply_single_movement(
            db, loc, data.movement_type, data.code, data.name, data.quantity, data.note, current_user.id,
            device_model=data.device_model, part_status=data.part_status, related_branch=data.related_branch,
        )
        db.commit()
        return {
            "status": "success",
            "code": part.code,
            "name": part.name,
            "location": loc,
            "movement_type": data.movement_type,
            "new_quantity": stock.quantity,
        }
    except ValueError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        db.rollback()
        logger.error(f"Error saving stock movement: {str(e)}")
        raise HTTPException(status_code=400, detail=f"Gagal menyimpan pergerakan stok: {str(e)}")


@router.get("/movement/template")
def download_movement_template(
    kind: str = "terima",
    current_user: User = Depends(get_current_user),
):
    """
    Contoh format Excel untuk upload massal pergerakan stok - file ASLI yang
    diberikan (bukan digenerate ulang), disimpan di app/assets/. Ada 2 versi:
    - kind=terima -> dipakai utk Terima dari Gudang / Terima dari Cabang / Terima dari Pusat
    - kind=kirim  -> dipakai utk Kirim ke Cabang / Kirim Balik ke Pusat
    """
    import os
    filename_map = {
        "terima": ("Contoh_Format_Upload_Terima_Stok.xlsx", "Contoh_Format_Upload_Terima_Stok.xlsx"),
        "kirim": ("Contoh_Format_Upload_Kirim_Stok.xlsx", "Contoh_Format_Upload_Kirim_Stok.xlsx"),
    }
    if kind not in filename_map:
        raise HTTPException(status_code=400, detail="kind harus 'terima' atau 'kirim'.")

    filename, download_name = filename_map[kind]
    # inventory_parts.py ada di app/api/v1/, jadi perlu naik DUA level (v1 -> api -> app)
    # baru masuk ke app/assets/ - beda dengan service_report_template.py yang ada
    # langsung di app/services/ (cukup naik satu level).
    file_path = os.path.join(os.path.dirname(__file__), "..", "..", "assets", filename)
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="File contoh format tidak ditemukan di server.")

    with open(file_path, "rb") as f:
        content = f.read()

    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={download_name}"},
    )


@router.post("/movement/bulk-upload")
def bulk_upload_movements(
    location: str = Form(...),
    movement_type: str = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Upload massal pergerakan stok dari Excel, format kolom (7 kolom, sesuai
    file "Contoh_Format_Upload_Terima_Stok.xlsx" / "..._Kirim_Stok.xlsx"):
    Nama Sparepart, Kode Sparepart, Status, Model Alat, Jumlah,
    Nama Cabang Asal/Tujuan, Catatan.
    (baris header terdeteksi otomatis & dilewati).
    Setiap baris diproses dengan aturan yang SAMA PERSIS dengan input manual
    satu-satu (stok tidak boleh minus, Cabang Asal/Tujuan wajib utk jenis
    pergerakan tertentu, dst) - baris yang gagal dilaporkan jelas tanpa
    membatalkan baris lain yang valid.
    """
    loc = location.lower().strip()
    if loc not in VALID_INV_LOCATIONS:
        raise HTTPException(status_code=400, detail=f"location harus salah satu dari: {', '.join(VALID_INV_LOCATIONS)}.")
    require_department_access(current_user, loc)

    allowed = MOVEMENT_RULES.get(loc, {})
    if movement_type not in allowed:
        raise HTTPException(
            status_code=400,
            detail=f"movement_type '{movement_type}' tidak berlaku untuk lokasi '{loc}'. "
                   f"Pilihan yang valid: {', '.join(allowed.keys())}.",
        )

    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="File harus berformat .xlsx (Excel).")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(file.file.read()), data_only=True)
        ws = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca file Excel: {str(e)}")

    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="File Excel kosong.")

    header_probe = " ".join(str(v).lower() for v in rows[0] if v)
    start_row = 1 if ("sparepart" in header_probe or "kode" in header_probe or "jumlah" in header_probe) else 0

    # ---- Ambil SEMUA data yang mungkin dibutuhkan di awal, sekali saja ----
    # (bukan query berulang di dalam loop per baris) - untuk file berisi ratusan
    # baris, ini memangkas ratusan query database jadi hanya beberapa. Ini
    # penyebab utama proses lambat/timeout ("upstream error") pada file besar.
    branch_cache = {b.name.lower(): b.name for b in db.query(BranchCatalog).filter(BranchCatalog.is_active.is_(True)).all()}

    distinct_codes = {
        str(row[1]).strip() for row in rows[start_row:]
        if row and len(row) > 1 and row[1] is not None and str(row[1]).strip()
    }
    part_cache = {p.code: p for p in db.query(PartCatalog).filter(PartCatalog.code.in_(distinct_codes)).all()} if distinct_codes else {}

    stock_cache = {}
    if part_cache:
        existing_part_ids = [p.id for p in part_cache.values()]
        for s in db.query(PartStock).filter(PartStock.part_id.in_(existing_part_ids)).with_for_update().all():
            stock_cache[(s.part_id, s.location, s.branch_name)] = s

    def _cell(row, i):
        return str(row[i]).strip() if len(row) > i and row[i] is not None else None

    berhasil, errors = 0, []
    for idx, row in enumerate(rows[start_row:], start=start_row + 1):
        if not row or not any(row):
            continue
        name = _cell(row, 0)             # Nama Sparepart
        code = _cell(row, 1) or ""       # Kode Sparepart
        part_status = _cell(row, 2)      # Status
        device_model = _cell(row, 3)     # Model Alat
        try:
            qty_raw = row[4] if len(row) > 4 else None
            quantity = int(qty_raw) if qty_raw is not None else 0
        except (ValueError, TypeError):
            quantity = 0
        related_branch = _cell(row, 5)   # Nama Cabang Asal/Tujuan
        note = _cell(row, 6)             # Catatan

        try:
            _apply_single_movement(
                db, loc, movement_type, code, name, quantity, note, current_user.id,
                device_model=device_model, part_status=part_status, related_branch=related_branch,
                branch_cache=branch_cache, part_cache=part_cache, stock_cache=stock_cache,
            )
            berhasil += 1
        except ValueError as e:
            errors.append({"row": idx, "reason": str(e)})
        except Exception as e:
            errors.append({"row": idx, "reason": f"Error tidak terduga: {str(e)}"})

    db.commit()
    return {
        "status": "success",
        "total_baris_diproses": len(rows) - start_row,
        "berhasil": berhasil,
        "gagal": len(errors),
        "detail_gagal": errors[:50],
    }



def _apply_single_opname(
    db: Session, loc: str, code: str, name: Optional[str], counted_quantity: int,
    note: Optional[str], user_id: int, branch_name: Optional[str] = None,
    branch_cache: Optional[dict] = None, part_cache: Optional[dict] = None, stock_cache: Optional[dict] = None,
):
    """
    Logika inti SATU baris stok opname - dipakai bersama oleh endpoint manual
    (/opname) MAUPUN upload massal Excel (/opname/bulk-upload), sama seperti
    pola _apply_single_movement.
    """
    if not code or not code.strip():
        raise ValueError("Kode Sparepart wajib diisi.")
    if counted_quantity is None or counted_quantity < 0:
        raise ValueError("Jumlah Hasil Hitung Fisik wajib diisi (angka >= 0).")

    clean_branch = (branch_name or "").strip()
    if loc == "cabang":
        if not clean_branch:
            raise ValueError("Nama Cabang wajib diisi untuk stok opname di Cabang.")
        branch_key = clean_branch.lower()
        if branch_cache is not None:
            official_name = branch_cache.get(branch_key)
        else:
            branch_match = (
                db.query(BranchCatalog)
                .filter(BranchCatalog.name.ilike(clean_branch), BranchCatalog.is_active.is_(True))
                .first()
            )
            official_name = branch_match.name if branch_match else None
        if not official_name:
            raise ValueError(f"Nama cabang '{clean_branch}' tidak ditemukan di daftar Kelola Cabang (atau sedang nonaktif).")
        clean_branch = official_name

    part = _get_or_create_part(db, code.strip(), name, part_cache=part_cache)
    stock = _get_or_create_stock_row(db, part.id, loc, branch_name=clean_branch, stock_cache=stock_cache)

    system_qty = stock.quantity
    difference = counted_quantity - system_qty

    db.add(PartStockOpname(
        part_id=part.id,
        location=loc,
        branch_name=clean_branch or None,
        system_quantity=system_qty,
        counted_quantity=counted_quantity,
        difference=difference,
        note=note,
        performed_by_user_id=user_id,
    ))
    stock.quantity = counted_quantity
    return part, stock, system_qty, difference


@router.post("/opname")
def create_opname(
    data: OpnameIn,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    loc = data.location.lower().strip()
    if loc not in VALID_INV_LOCATIONS:
        raise HTTPException(status_code=400, detail=f"location harus salah satu dari: {', '.join(VALID_INV_LOCATIONS)}.")
    require_department_access(current_user, loc)

    try:
        part, stock, system_qty, difference = _apply_single_opname(
            db, loc, data.code, data.name, data.counted_quantity, data.note, current_user.id,
            branch_name=data.branch_name,
        )
        db.commit()
        return {
            "status": "success",
            "code": part.code,
            "name": part.name,
            "location": loc,
            "system_quantity": system_qty,
            "counted_quantity": data.counted_quantity,
            "difference": difference,
        }
    except ValueError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        db.rollback()
        logger.error(f"Error saving stock opname: {str(e)}")
        raise HTTPException(status_code=400, detail=f"Gagal menyimpan stok opname: {str(e)}")


@router.get("/opname/template")
def download_opname_template(
    location: str = "pusat",
    current_user: User = Depends(get_current_user),
):
    """Contoh format Excel utk upload massal Stok Opname - dibuat langsung (bukan file statis)."""
    loc = location.lower().strip()
    headers = ["Nama Sparepart", "Kode Sparepart", "Jumlah Hasil Hitung Fisik", "Catatan"]
    sample = ["Contoh: LCD Panel", "LCD-001", 25, "Opsional"]
    if loc == "cabang":
        headers.insert(2, "Nama Cabang")
        sample.insert(2, "Contoh: Cabang Medan")
    excel_file = generate_inventory_excel(headers, [sample], "Stok Opname")
    filename = f"Contoh_Format_Upload_Stok_Opname_{loc}.xlsx"
    return StreamingResponse(
        excel_file,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.post("/opname/bulk-upload")
def bulk_upload_opname(
    location: str = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Upload massal Stok Opname dari Excel. Format kolom:
    - Pusat: Nama Sparepart, Kode Sparepart, Jumlah Hasil Hitung Fisik, Catatan (4 kolom)
    - Cabang: Nama Sparepart, Kode Sparepart, Nama Cabang, Jumlah Hasil Hitung Fisik, Catatan (5 kolom)
    Baris header terdeteksi otomatis & dilewati. Baris yang gagal dilaporkan
    jelas tanpa membatalkan baris lain yang valid (sama seperti bulk-upload pergerakan).
    """
    loc = location.lower().strip()
    if loc not in VALID_INV_LOCATIONS:
        raise HTTPException(status_code=400, detail=f"location harus salah satu dari: {', '.join(VALID_INV_LOCATIONS)}.")
    require_department_access(current_user, loc)

    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="File harus berformat .xlsx (Excel).")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(file.file.read()), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca file Excel: {str(e)}")

    if not rows:
        raise HTTPException(status_code=400, detail="File Excel kosong.")

    start_row = 1 if rows and rows[0] and str(rows[0][0]).strip().lower() in ("nama sparepart", "kode sparepart") else 0

    branch_cache = {b.name.lower(): b.name for b in db.query(BranchCatalog).filter(BranchCatalog.is_active.is_(True)).all()}
    distinct_codes = {
        str(row[1]).strip() for row in rows[start_row:]
        if row and len(row) > 1 and row[1] is not None and str(row[1]).strip()
    }
    part_cache = {p.code: p for p in db.query(PartCatalog).filter(PartCatalog.code.in_(distinct_codes)).all()} if distinct_codes else {}
    stock_cache = {}
    if part_cache:
        existing_part_ids = [p.id for p in part_cache.values()]
        for s in db.query(PartStock).filter(PartStock.part_id.in_(existing_part_ids)).with_for_update().all():
            stock_cache[(s.part_id, s.location, s.branch_name)] = s

    def _cell(row, i):
        return str(row[i]).strip() if len(row) > i and row[i] is not None else None

    berhasil, errors = 0, []
    for idx, row in enumerate(rows[start_row:], start=start_row + 1):
        if not row or not any(row):
            continue
        name = _cell(row, 0)
        code = _cell(row, 1) or ""
        if loc == "cabang":
            branch_name = _cell(row, 2)
            try:
                qty_raw = row[3] if len(row) > 3 else None
                counted_quantity = int(qty_raw) if qty_raw is not None else None
            except (ValueError, TypeError):
                counted_quantity = None
            note = _cell(row, 4)
        else:
            branch_name = None
            try:
                qty_raw = row[2] if len(row) > 2 else None
                counted_quantity = int(qty_raw) if qty_raw is not None else None
            except (ValueError, TypeError):
                counted_quantity = None
            note = _cell(row, 3)

        try:
            _apply_single_opname(
                db, loc, code, name, counted_quantity, note, current_user.id,
                branch_name=branch_name, branch_cache=branch_cache, part_cache=part_cache, stock_cache=stock_cache,
            )
            berhasil += 1
        except ValueError as e:
            errors.append({"row": idx, "reason": str(e)})
        except Exception as e:
            errors.append({"row": idx, "reason": f"Error tidak terduga: {str(e)}"})

    db.commit()
    return {
        "status": "success",
        "total_baris_diproses": len(rows) - start_row,
        "berhasil": berhasil,
        "gagal": len(errors),
        "detail_gagal": errors[:50],
    }


# ---------------- Laporan Excel ----------------

def _excel_response(headers, rows, sheet_title, filename):
    excel_file = generate_inventory_excel(headers, rows, sheet_title)
    return StreamingResponse(
        excel_file,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/report/stock-list")
def report_stock_list(
    location: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    loc = location.lower().strip()
    require_department_access(current_user, loc)
    rows = (
        db.query(PartStock, PartCatalog)
        .join(PartCatalog, PartStock.part_id == PartCatalog.id)
        .filter(PartStock.location == loc)
        .order_by(PartCatalog.name, PartStock.branch_name)
        .all()
    )
    if loc == "cabang":
        headers = ["Kode Sparepart", "Nama Sparepart", "Nama Cabang", "Jumlah Stok", "Harga Satuan"]
        data = [[p.code, p.name or "-", s.branch_name or "-", s.quantity, p.unit_price or 0] for s, p in rows]
    else:
        headers = ["Kode Sparepart", "Nama Sparepart", "Jumlah Stok", "Harga Satuan"]
        data = [[p.code, p.name or "-", s.quantity, p.unit_price or 0] for s, p in rows]
    filename = f"Stok_Sparepart_{loc}_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return _excel_response(headers, data, f"Stok {loc}", filename)


@router.get("/report/movements")
def report_movements(
    location: str,
    movement_type: str,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    loc = location.lower().strip()
    require_department_access(current_user, loc)
    if movement_type not in REPORTABLE_MOVEMENT_TYPES.get(loc, {}):
        raise HTTPException(status_code=400, detail="movement_type tidak valid untuk lokasi ini.")

    query = (
        db.query(PartStockMovement, PartCatalog)
        .join(PartCatalog, PartStockMovement.part_id == PartCatalog.id)
        .filter(PartStockMovement.location == loc, PartStockMovement.movement_type == movement_type)
    )
    if date_from:
        query = query.filter(PartStockMovement.created_at >= datetime.fromisoformat(date_from))
    if date_to:
        query = query.filter(PartStockMovement.created_at < datetime.fromisoformat(date_to) + timedelta(days=1))
    rows = query.order_by(desc(PartStockMovement.created_at)).all()
    data = [
        [
            m.created_at.strftime("%Y-%m-%d %H:%M") if m.created_at else "-",
            p.code, p.name or "-", m.device_model or "-", m.part_status or "-",
            m.quantity, m.related_branch or "-", m.note or "-",
        ]
        for m, p in rows
    ]
    label = MOVEMENT_LABELS.get(movement_type, movement_type)
    filename = f"Laporan_{label.replace(' ', '_')}_{loc}_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return _excel_response(
        ["Tanggal", "Kode Sparepart", "Nama Sparepart", "Model Alat", "Status", "Jumlah", "Cabang Asal/Tujuan", "Catatan"],
        data, label[:31], filename,
    )


@router.get("/report/opname")
def report_opname(
    location: str,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    loc = location.lower().strip()
    require_department_access(current_user, loc)
    query = (
        db.query(PartStockOpname, PartCatalog)
        .join(PartCatalog, PartStockOpname.part_id == PartCatalog.id)
        .filter(PartStockOpname.location == loc)
    )
    if date_from:
        query = query.filter(PartStockOpname.created_at >= datetime.fromisoformat(date_from))
    if date_to:
        query = query.filter(PartStockOpname.created_at < datetime.fromisoformat(date_to) + timedelta(days=1))
    rows = query.order_by(desc(PartStockOpname.created_at)).all()
    data = [
        [
            o.created_at.strftime("%Y-%m-%d %H:%M") if o.created_at else "-",
            p.code, p.name or "-", o.system_quantity, o.counted_quantity, o.difference, o.note or "-",
        ]
        for o, p in rows
    ]
    filename = f"Stok_Opname_{loc}_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return _excel_response(
        ["Tanggal", "Kode Sparepart", "Nama Sparepart", "Stok Sistem", "Hasil Hitung Fisik", "Selisih", "Catatan"],
        data, f"Opname {loc}", filename,
    )
