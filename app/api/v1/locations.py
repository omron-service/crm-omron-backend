import logging
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from pydantic import BaseModel, Field

from app.db.session import get_db
from app.core.deps import get_current_user, require_role
from app.models.schema import BranchCatalog, PickupCenterCatalog, User

logger = logging.getLogger("uvicorn.error")

branch_router = APIRouter(prefix="/api/v1/branches", tags=["Kelola Cabang"])
pickup_router = APIRouter(prefix="/api/v1/pickup-centers", tags=["Kelola Pickup Center"])


# ==================== Skema ====================

class BranchIn(BaseModel):
    name: str = Field(max_length=150)
    code: str = Field(max_length=20)
    handled_by: Optional[str] = Field(default=None, max_length=150)
    city: Optional[str] = Field(default=None, max_length=100)
    address: Optional[str] = None


class PickupCenterIn(BaseModel):
    store_long_code: str = Field(max_length=50)
    store_name: str = Field(max_length=200)
    store_address: Optional[str] = None
    city: Optional[str] = Field(default=None, max_length=100)


# ==================== CABANG ====================

@branch_router.get("/")
def list_branches(
    include_inactive: bool = False,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(BranchCatalog)
    if not include_inactive:
        query = query.filter(BranchCatalog.is_active.is_(True))
    return query.order_by(BranchCatalog.name).all()


@branch_router.post("/", status_code=201)
def create_branch(
    data: BranchIn,
    db: Session = Depends(get_db),
    _admin: User = Depends(require_role("superadmin")),
):
    code = data.code.strip().upper()
    if db.query(BranchCatalog).filter(BranchCatalog.code == code).first():
        raise HTTPException(status_code=400, detail=f"Kode Store '{code}' sudah dipakai cabang lain.")
    entry = BranchCatalog(
        name=data.name.strip(), code=code, handled_by=data.handled_by,
        city=data.city, address=data.address,
    )
    db.add(entry)
    db.commit()
    db.refresh(entry)
    return entry


@branch_router.put("/{branch_id}")
def update_branch(
    branch_id: int,
    data: BranchIn,
    db: Session = Depends(get_db),
    _admin: User = Depends(require_role("superadmin")),
):
    entry = db.query(BranchCatalog).filter(BranchCatalog.id == branch_id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Cabang tidak ditemukan.")
    new_code = data.code.strip().upper()
    if new_code != entry.code and db.query(BranchCatalog).filter(BranchCatalog.code == new_code).first():
        raise HTTPException(status_code=400, detail=f"Kode Store '{new_code}' sudah dipakai cabang lain.")
    entry.name = data.name.strip()
    entry.code = new_code
    entry.handled_by = data.handled_by
    entry.city = data.city
    entry.address = data.address
    db.commit()
    db.refresh(entry)
    return entry


@branch_router.patch("/{branch_id}/deactivate")
def deactivate_branch(
    branch_id: int, db: Session = Depends(get_db), _admin: User = Depends(require_role("superadmin")),
):
    entry = db.query(BranchCatalog).filter(BranchCatalog.id == branch_id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Cabang tidak ditemukan.")
    entry.is_active = False
    db.commit()
    return {"status": "success"}


@branch_router.patch("/{branch_id}/reactivate")
def reactivate_branch(
    branch_id: int, db: Session = Depends(get_db), _admin: User = Depends(require_role("superadmin")),
):
    entry = db.query(BranchCatalog).filter(BranchCatalog.id == branch_id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Cabang tidak ditemukan.")
    entry.is_active = True
    db.commit()
    return {"status": "success"}


@branch_router.post("/bulk-upload")
def bulk_upload_branches(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _admin: User = Depends(require_role("superadmin")),
):
    """
    Upload dari Excel, kolom: Nama Cabang, Kode Store, Handle by, Kota, Alamat
    (urutan kolom sesuai template; baris header terdeteksi otomatis & dilewati).

    Kode Store yang SUDAH ADA -> data lain (nama, handle by, kota, alamat)
    DIPERBARUI mengikuti isi file (dan diaktifkan lagi kalau sebelumnya nonaktif).
    Kode Store yang BELUM ADA -> ditambahkan sebagai cabang baru.
    Cabang yang tidak disebut di file TIDAK disentuh (tidak dinonaktifkan/dihapus).
    """
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="File harus berformat .xlsx (Excel).")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(file.file.read()), data_only=True)
        ws = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca file Excel: {str(e)}")

    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="File Excel kosong.")

    header_probe = " ".join(str(v).lower() for v in rows[0] if v)
    start_row = 1 if ("kode" in header_probe or "cabang" in header_probe or "store" in header_probe) else 0

    created, updated, errors = 0, 0, []
    for idx, row in enumerate(rows[start_row:], start=start_row + 1):
        if not row or not any(row):
            continue
        name = str(row[0]).strip() if len(row) > 0 and row[0] else ""
        code = str(row[1]).strip().upper() if len(row) > 1 and row[1] else ""
        handled_by = str(row[2]).strip() if len(row) > 2 and row[2] else None
        city = str(row[3]).strip() if len(row) > 3 and row[3] else None
        address = str(row[4]).strip() if len(row) > 4 and row[4] else None

        if not name or not code:
            errors.append({"row": idx, "reason": "Nama Cabang atau Kode Store kosong"})
            continue

        existing = db.query(BranchCatalog).filter(BranchCatalog.code == code).first()
        if existing:
            existing.name = name
            existing.handled_by = handled_by
            existing.city = city
            existing.address = address
            existing.is_active = True
            updated += 1
        else:
            db.add(BranchCatalog(name=name, code=code, handled_by=handled_by, city=city, address=address))
            created += 1

    db.commit()
    return {
        "status": "success",
        "total_baris_diproses": len(rows) - start_row,
        "ditambahkan_baru": created,
        "diperbarui": updated,
        "gagal": len(errors),
        "detail_gagal": errors[:50],
    }


# ==================== PICKUP CENTER ====================

@pickup_router.get("/")
def list_pickup_centers(
    include_inactive: bool = False,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(PickupCenterCatalog)
    if not include_inactive:
        query = query.filter(PickupCenterCatalog.is_active.is_(True))
    return query.order_by(PickupCenterCatalog.store_name).all()


@pickup_router.post("/", status_code=201)
def create_pickup_center(
    data: PickupCenterIn,
    db: Session = Depends(get_db),
    _admin: User = Depends(require_role("superadmin")),
):
    code = data.store_long_code.strip()
    if db.query(PickupCenterCatalog).filter(PickupCenterCatalog.store_long_code == code).first():
        raise HTTPException(status_code=400, detail=f"Store Long Code '{code}' sudah dipakai.")
    entry = PickupCenterCatalog(
        store_long_code=code, store_name=data.store_name.strip(),
        store_address=data.store_address, city=data.city,
    )
    db.add(entry)
    db.commit()
    db.refresh(entry)
    return entry


@pickup_router.put("/{pickup_id}")
def update_pickup_center(
    pickup_id: int,
    data: PickupCenterIn,
    db: Session = Depends(get_db),
    _admin: User = Depends(require_role("superadmin")),
):
    entry = db.query(PickupCenterCatalog).filter(PickupCenterCatalog.id == pickup_id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Pickup Center tidak ditemukan.")
    new_code = data.store_long_code.strip()
    if new_code != entry.store_long_code and db.query(PickupCenterCatalog).filter(
        PickupCenterCatalog.store_long_code == new_code
    ).first():
        raise HTTPException(status_code=400, detail=f"Store Long Code '{new_code}' sudah dipakai.")
    entry.store_long_code = new_code
    entry.store_name = data.store_name.strip()
    entry.store_address = data.store_address
    entry.city = data.city
    db.commit()
    db.refresh(entry)
    return entry


@pickup_router.patch("/{pickup_id}/deactivate")
def deactivate_pickup_center(
    pickup_id: int, db: Session = Depends(get_db), _admin: User = Depends(require_role("superadmin")),
):
    entry = db.query(PickupCenterCatalog).filter(PickupCenterCatalog.id == pickup_id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Pickup Center tidak ditemukan.")
    entry.is_active = False
    db.commit()
    return {"status": "success"}


@pickup_router.patch("/{pickup_id}/reactivate")
def reactivate_pickup_center(
    pickup_id: int, db: Session = Depends(get_db), _admin: User = Depends(require_role("superadmin")),
):
    entry = db.query(PickupCenterCatalog).filter(PickupCenterCatalog.id == pickup_id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Pickup Center tidak ditemukan.")
    entry.is_active = True
    db.commit()
    return {"status": "success"}


@pickup_router.post("/bulk-upload")
def bulk_upload_pickup_centers(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _admin: User = Depends(require_role("superadmin")),
):
    """
    Upload dari Excel, kolom: Store Long Code, Store Name, Store Address, City.
    Aturan sama seperti upload Cabang: yang sudah ada DIPERBARUI datanya (dan
    diaktifkan lagi kalau nonaktif), yang belum ada DITAMBAHKAN, yang tidak
    disebut di file TIDAK disentuh sama sekali.
    """
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="File harus berformat .xlsx (Excel).")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(file.file.read()), data_only=True)
        ws = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca file Excel: {str(e)}")

    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="File Excel kosong.")

    header_probe = " ".join(str(v).lower() for v in rows[0] if v)
    start_row = 1 if ("store" in header_probe or "code" in header_probe) else 0

    created, updated, errors = 0, 0, []
    for idx, row in enumerate(rows[start_row:], start=start_row + 1):
        if not row or not any(row):
            continue
        code = str(row[0]).strip() if len(row) > 0 and row[0] else ""
        name = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        address = str(row[2]).strip() if len(row) > 2 and row[2] else None
        city = str(row[3]).strip() if len(row) > 3 and row[3] else None

        if not code or not name:
            errors.append({"row": idx, "reason": "Store Long Code atau Store Name kosong"})
            continue

        existing = db.query(PickupCenterCatalog).filter(PickupCenterCatalog.store_long_code == code).first()
        if existing:
            existing.store_name = name
            existing.store_address = address
            existing.city = city
            existing.is_active = True
            updated += 1
        else:
            db.add(PickupCenterCatalog(store_long_code=code, store_name=name, store_address=address, city=city))
            created += 1

    db.commit()
    return {
        "status": "success",
        "total_baris_diproses": len(rows) - start_row,
        "ditambahkan_baru": created,
        "diperbarui": updated,
        "gagal": len(errors),
        "detail_gagal": errors[:50],
    }
